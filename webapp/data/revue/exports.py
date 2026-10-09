"""Excel exports of the review screen."""

import re
from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter

from data.models.suggestion import SuggestionCohorte
from data.revue.cohortes import logs_queryset

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
# Excel refuses longer cells
MAX_CELL_LENGTH = 32_767

LOG_COLUMNS = [
    ("Niveau", 10),
    ("Date", 18),
    ("Message", 60),
    ("Fonction de transformation", 30),
    ("Identifiant unique", 30),
    ("Groupe", 10),
    ("Colonnes d'origine", 30),
    ("Valeurs d'origine", 40),
    ("Colonnes de destination", 30),
]


def _cell(value):
    if value is None:
        return None
    if isinstance(value, list):
        value = ", ".join(str(item) for item in value)
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)[:MAX_CELL_LENGTH]
    return value


def export_filename(cohorte: SuggestionCohorte) -> str:
    action = re.sub(r"[^A-Za-z0-9_-]+", "_", cohorte.identifiant_action).strip("_")
    return f"logs_cohorte_{cohorte.id}_{action}.xlsx"


def logs_xlsx(cohorte: SuggestionCohorte, niveau: list[str]) -> bytes:
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet(f"Logs cohorte {cohorte.id}")
    for index, (_, width) in enumerate(LOG_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.append([title for title, _ in LOG_COLUMNS])
    for log in logs_queryset(cohorte, niveau).iterator(chunk_size=2000):
        sheet.append(
            [
                _cell(log.niveau_de_log),
                # Excel has no timezone: local time of the server
                timezone.localtime(log.cree_le).replace(tzinfo=None),
                _cell(log.message),
                _cell(log.fonction_de_transformation),
                _cell(log.identifiant_unique),
                log.suggestion_groupe_id,
                _cell(log.origine_colonnes),
                _cell(log.origine_valeurs),
                _cell(log.destination_colonnes),
            ]
        )
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
